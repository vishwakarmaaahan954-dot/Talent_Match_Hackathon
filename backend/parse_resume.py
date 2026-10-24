"""
RADIX Talent Match Hackathon - Role 2: Resume Parsing

Usage (from the command line):
    python parse_resume.py samples/anny.pdf
    python parse_resume.py samples/anny.docx --no-llm

What it does:
    1. Reads text from a PDF / DOCX / TXT resume
    2. Extracts skills tagged with one of the 12 RADIX category codes
       - with an LLM (Claude) if ANTHROPIC_API_KEY is set
       - otherwise with a keyword dictionary (works with no API key)
    3. Pulls name, email, education, projects, experience, certifications
    4. Returns / saves JSON in the shared data contract format
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path

# ---------------------------------------------------------------------------
# 1. The shared data contract: 12 categories + OTHER
# ---------------------------------------------------------------------------
CATEGORIES = {
    "COD": "Coding (programming languages, writing code, frameworks)",
    "DSA": "Data Structures & Algorithms",
    "OOD": "Object-Oriented Design",
    "APTI": "Aptitude (quantitative, logical reasoning)",
    "COMM": "Communication (speaking, writing, teamwork, leadership)",
    "AI": "AI (machine learning, deep learning, NLP, LLMs, data science)",
    "CLOUD": "Cloud (AWS, Azure, GCP, Docker, Kubernetes, DevOps)",
    "SQL": "SQL (databases and queries)",
    "SWE": "Software Engineering (git, testing, agile, APIs, SDLC)",
    "SYSD": "System Design (scalability, architecture, distributed systems)",
    "NETW": "Networking",
    "OS": "Operating Systems (Linux, processes, threads)",
}
VALID_CODES = set(CATEGORIES) | {"OTHER"}
VALID_CONFIDENCE = {"high", "medium", "low"}

# Keyword dictionary used when no LLM is available (the "fallback" mode)
KEYWORDS = {
    "COD": [
        "python", "java", "c++", "c#", "javascript", "typescript", "golang",
        "kotlin", "swift", "php", "ruby", "rust", "scala", "html", "css",
        "react", "angular", "vue", "node.js", "nodejs", "flask", "django",
        "fastapi", "spring boot", "spring", "matlab", "r programming",
    ],
    "DSA": [
        "data structures", "algorithms", "dsa", "leetcode", "codechef",
        "codeforces", "hackerrank", "dynamic programming", "graph algorithms",
        "sorting", "binary search", "competitive programming",
    ],
    "OOD": [
        "oop", "oops", "object-oriented", "object oriented", "design patterns",
        "solid principles", "uml", "class diagram",
    ],
    "APTI": [
        "aptitude", "quantitative", "logical reasoning", "analytical",
        "problem solving", "problem-solving", "puzzles", "cat percentile",
    ],
    "COMM": [
        "communication", "presentation", "public speaking", "teamwork",
        "team lead", "leadership", "mentoring", "collaboration",
        "technical writing", "debate",
    ],
    "AI": [
        "machine learning", "deep learning", "nlp", "natural language processing",
        "tensorflow", "pytorch", "keras", "scikit-learn", "sklearn", "pandas",
        "numpy", "llm", "langchain", "computer vision", "opencv", "data science",
        "generative ai", "neural network", "transformers", "data analysis",
        "power bi", "tableau", "statistics",
    ],
    "CLOUD": [
        "aws", "azure", "gcp", "google cloud", "docker", "kubernetes", "ci/cd",
        "terraform", "jenkins", "devops", "cloud computing", "lambda", "ec2",
    ],
    "SQL": [
        "sql", "mysql", "postgresql", "postgres", "oracle", "sqlite", "dbms",
        "database", "mongodb", "supabase", "pl/sql", "nosql",
    ],
    "SWE": [
        "git", "github", "agile", "scrum", "unit testing", "testing", "rest api",
        "rest apis", "restful", "microservices", "jira", "debugging", "sdlc",
        "code review", "software engineering", "api development", "junit", "pytest",
    ],
    "SYSD": [
        "system design", "scalability", "load balancing", "caching",
        "distributed systems", "high availability", "low level design",
        "high level design", "message queue", "kafka", "redis",
    ],
    "NETW": [
        "networking", "computer networks", "tcp/ip", "http", "dns", "osi model",
        "routing", "socket programming", "cisco", "ccna", "wireshark",
    ],
    "OS": [
        "operating system", "operating systems", "linux", "unix", "multithreading",
        "process scheduling", "shell scripting", "bash", "memory management",
        "deadlock", "ubuntu",
    ],
}

# ---------------------------------------------------------------------------
# 2. Read text out of the file
# ---------------------------------------------------------------------------
def extract_text(path: str) -> str:
    """Return the plain text of a PDF, DOCX or TXT file."""
    p = Path(path)
    ext = p.suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(p))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages)

    if ext == ".docx":
        from docx import Document
        doc = Document(str(p))
        lines = [para.text for para in doc.paragraphs]
        # many resumes keep skills inside tables, so read those too
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)

    if ext == ".txt":
        return p.read_text(encoding="utf-8", errors="ignore")

    raise ValueError(f"Unsupported file type: {ext} (use PDF, DOCX or TXT)")


# ---------------------------------------------------------------------------
# 3. Basic fields: name, email, sections
# ---------------------------------------------------------------------------
SECTION_HEADINGS = {
    "education": r"education|academic",
    "projects": r"projects?|academic projects?|personal projects?",
    "experience": r"(work |professional )?experience|employment",
    "internships": r"internships?|internship experience|training",
    "hackathons": r"hackathons?|competitions?|hackathons? ?(and|&) ?competitions?",
    "achievements": r"(achievements?|awards?|honou?rs|accomplishments?)"
                    r"( ?(and|&) ?(achievements?|awards?|honou?rs|accomplishments?))?",
    # combined headings like "Hackathons / Internships" -> sorted line by line later
    "mixed": r"(hackathons?|internships?|achievements?|awards?|competitions?)"
             r"( ?(/|and|&) ?(hackathons?|internships?|achievements?|awards?|competitions?))+",
    "certifications": r"certifications?|certificates?|courses",
    "skills": r"(technical )?skills|technologies|tech stack",
    # headings we don't use, but which must END the previous section
    "other": r"(preferred|target) roles?|job preferences?|(career )?objective|summary|"
             r"profile|languages?|interests?|hobbies|references?|declaration|personal details?",
}

HACKATHON_HINT = re.compile(r"hackathon|\bsih\b|smart india|hackfest|codeathon|ideathon", re.I)
INTERN_HINT = re.compile(r"\bintern(ship)?s?\b|trainee", re.I)


def _heading_of(line: str):
    clean = re.sub(r"[^A-Za-z /&]", "", line).strip().lower()
    if not clean or len(clean) > 40:
        return None
    for key, pattern in SECTION_HEADINGS.items():
        if re.fullmatch(pattern, clean):
            return key
    return None


def split_sections(text: str) -> dict:
    sections, current = {}, None
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        head = _heading_of(line)
        if head:
            current = head
            sections.setdefault(current, [])
        elif current:
            entry = line.lstrip("•-*▪● ").strip()
            # PDF line-wrap: a lowercase start means the previous entry continues
            if current not in ("skills", "other") and sections[current] and entry[:1].islower():
                sections[current][-1] += " " + entry
            else:
                sections[current].append(entry)
    return sections


def split_extra_sections(sections: dict):
    """Return (internships, hackathons, achievements); sorts combined sections by content."""
    intern = list(sections.get("internships", []))
    hack = list(sections.get("hackathons", []))
    ach = []
    pool = [(l, True) for l in sections.get("mixed", [])] + \
           [(l, False) for l in sections.get("achievements", [])]
    for line, from_mixed in pool:
        if HACKATHON_HINT.search(line):
            hack.append(line)
        elif from_mixed and INTERN_HINT.search(line):
            intern.append(line)
        else:
            ach.append(line)
    return intern, hack, ach


def extract_basics(text: str) -> dict:
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    name = ""
    for line in text.splitlines():
        line = line.strip()
        if line and "@" not in line and len(line) < 50 and not re.search(r"\d{5,}", line):
            name = line.title() if line.isupper() else line
            break
    return {"name": name, "email": email.group(0) if email else ""}


# ---------------------------------------------------------------------------
# 4a. Skill extraction WITHOUT an LLM (keyword matching)
# ---------------------------------------------------------------------------
ACRONYMS = {
    "sql", "nlp", "aws", "gcp", "dsa", "oop", "oops", "uml", "dns", "http", "llm",
    "html", "css", "php", "dbms", "sdlc", "ci/cd", "tcp/ip", "pl/sql", "nosql",
    "ccna", "osi model", "rest api", "rest apis",
}
SPECIAL_NAMES = {
    "javascript": "JavaScript", "typescript": "TypeScript", "node.js": "Node.js",
    "nodejs": "Node.js", "mysql": "MySQL", "postgresql": "PostgreSQL",
    "mongodb": "MongoDB", "github": "GitHub", "tensorflow": "TensorFlow",
    "pytorch": "PyTorch", "scikit-learn": "scikit-learn", "sklearn": "scikit-learn",
    "numpy": "NumPy", "fastapi": "FastAPI", "langchain": "LangChain",
    "leetcode": "LeetCode", "hackerrank": "HackerRank", "codechef": "CodeChef",
    "codeforces": "Codeforces", "power bi": "Power BI", "opencv": "OpenCV",
    "golang": "Go", "supabase": "Supabase",
}


def pretty_name(word: str) -> str:
    if word in SPECIAL_NAMES:
        return SPECIAL_NAMES[word]
    if word in ACRONYMS:
        return word.upper().replace("REST APIS", "REST APIs")
    return word.title()


def _norm(line: str) -> str:
    return line.strip().lstrip("•-*▪● ").strip()


def _snippet(line: str, match) -> str:
    """Short piece of the line around the keyword (max ~160 characters)."""
    line = " ".join(line.split())
    if len(line) <= 160:
        return line
    start = max(0, match.start() - 70)
    end = min(len(line), match.end() + 80)
    return ("..." if start > 0 else "") + line[start:end].strip() + ("..." if end < len(line) else "")


def keyword_extract(text: str, sections: dict) -> list:
    lines = [_norm(l) for l in text.splitlines() if l.strip()]
    skills_lines = {_norm(l) for l in sections.get("skills", [])}

    def is_list_line(line: str) -> bool:
        # a skills section line, or a comma / pipe separated list of skills
        return line in skills_lines or line.count(",") >= 3 or line.count("|") >= 2

    found = []
    for code, words in KEYWORDS.items():
        for word in words:
            pattern = r"(?<![A-Za-z0-9+#])" + re.escape(word) + r"(?![A-Za-z0-9+#])"
            used = []    # lines in projects / experience that mention the skill
            listed = []  # lines that are just skill lists
            for line in lines:
                m = re.search(pattern, line, flags=re.IGNORECASE)
                if m:
                    (listed if is_list_line(line) else used).append((line, m))
            if not used and not listed:
                continue

            if used and listed:
                confidence = "high"      # listed AND used in a project / job
            elif used:
                confidence = "medium"    # used somewhere
            else:
                confidence = "low"       # only listed in the skills section

            # evidence: a sentence where the skill is actually used, if there is one
            evidence = _snippet(*used[0]) if used else "Listed in the skills section"
            found.append({
                "skill_name": pretty_name(word),
                "category_code": code,
                "evidence": evidence,
                "confidence": confidence,
            })
    return found


# ---------------------------------------------------------------------------
# 4b. Skill extraction WITH an LLM (Claude) - better at messy resumes
# ---------------------------------------------------------------------------
LLM_PROMPT = """You are a resume analyst. Read the resume below and extract the candidate's technical skills.

Tag every skill with EXACTLY ONE category_code from this list:
{categories}
Use "OTHER" for anything that fits none of them.

Rules:
- Include skills listed explicitly AND skills clearly shown in projects, internships or experience
  (e.g. "built a REST API with Flask and PostgreSQL" -> Flask (COD), REST APIs (SWE), PostgreSQL (SQL)).
- Do NOT invent skills that the resume does not show.
- "evidence" = a short quote from the resume or a short reason (max 20 words).
- "confidence": "high" = used in a project/job AND listed, "medium" = used or clearly implied, "low" = only listed.
- Also extract education, projects, experience, internships, hackathons, achievements and certifications as short strings.
- "internships" = internships/trainee roles only; "experience" = other jobs. Do not repeat an item in both.
- "hackathons" = hackathon name + project/result/position (e.g. "Smart India Hackathon 2025 - built X, finalist").
- "achievements" = awards, ranks, scholarships, competition wins that are not hackathons.

Return ONLY valid JSON, no other text, in exactly this shape:
{{
  "name": "...",
  "email": "...",
  "education": "...",
  "projects": ["..."],
  "experience": ["..."],
  "internships": ["..."],
  "hackathons": ["..."],
  "achievements": ["..."],
  "certifications": ["..."],
  "skills": [
    {{"skill_name": "...", "category_code": "...", "evidence": "...", "confidence": "high|medium|low"}}
  ]
}}

RESUME:
\"\"\"
{resume}
\"\"\"
"""


def llm_extract(text: str) -> dict:
    import anthropic  # imported here so the keyword mode works without it

    client = anthropic.Anthropic()  # reads ANTHROPIC_API_KEY from the environment
    model = os.environ.get("ANTHROPIC_MODEL", "claude-sonnet-5-5")
    categories = "\n".join(f"- {c}: {d}" for c, d in CATEGORIES.items())
    prompt = LLM_PROMPT.format(categories=categories, resume=text[:20000])

    last_error = None
    for _ in range(2):  # one retry if the JSON comes back broken
        reply = client.messages.create(
            model=model,
            max_tokens=4000,
            temperature=0,
            messages=[{"role": "user", "content": prompt}],
        )
        raw = reply.content[0].text
        try:
            return json.loads(raw[raw.index("{"): raw.rindex("}") + 1])
        except (ValueError, json.JSONDecodeError) as e:
            last_error = e
    raise RuntimeError(f"LLM did not return valid JSON: {last_error}")


# ---------------------------------------------------------------------------
# 5. Clean up so the output always matches the contract
# ---------------------------------------------------------------------------
def clean_skills(skills: list) -> list:
    seen, cleaned = set(), []
    for s in skills:
        name = str(s.get("skill_name", "")).strip()
        if not name or name.lower() in seen:
            continue
        seen.add(name.lower())
        code = str(s.get("category_code", "OTHER")).strip().upper()
        conf = str(s.get("confidence", "medium")).strip().lower()
        cleaned.append({
            "skill_name": name,
            "category_code": code if code in VALID_CODES else "OTHER",
            "evidence": str(s.get("evidence", "")).strip(),
            "confidence": conf if conf in VALID_CONFIDENCE else "medium",
        })
    return cleaned


# ---------------------------------------------------------------------------
# 6. The one function everybody else calls
# ---------------------------------------------------------------------------
def parse_resume(path: str, use_llm=None, save=True) -> dict:
    """
    Parse one resume file and return the shared-contract dict.
    use_llm=None  -> use the LLM only if ANTHROPIC_API_KEY is set
    use_llm=False -> keyword mode (no API key needed)
    """
    text = extract_text(path)
    if len(text.strip()) < 30:
        raise ValueError("Could not read any text from this file (scanned image PDF?).")

    sections = split_sections(text)
    basics = extract_basics(text)
    intern_fb, hack_fb, ach_fb = split_extra_sections(sections)
    if use_llm is None:
        use_llm = bool(os.environ.get("ANTHROPIC_API_KEY"))

    method = "keyword"
    llm_data = {}
    if use_llm:
        try:
            llm_data = llm_extract(text)
            method = "llm"
        except Exception as e:  # never crash the demo: fall back to keywords
            print(f"[warning] LLM failed ({e}); using keyword mode instead.", file=sys.stderr)

    if method == "llm":
        skills = clean_skills(llm_data.get("skills", []))
    else:
        skills = clean_skills(keyword_extract(text, sections))

    result = {
        "source_type": "resume",
        "source_file": Path(path).name,
        "company": "",
        "role": "",
        "name": llm_data.get("name") or basics["name"],
        "email": llm_data.get("email") or basics["email"],
        "education": llm_data.get("education") or "; ".join(sections.get("education", []))[:500],
        "projects": llm_data.get("projects") or sections.get("projects", [])[:10],
        "experience": llm_data.get("experience") or sections.get("experience", [])[:10],
        "internships": llm_data.get("internships") or intern_fb[:10],
        "hackathons": llm_data.get("hackathons") or hack_fb[:10],
        "achievements": llm_data.get("achievements") or ach_fb[:10],
        "certifications": llm_data.get("certifications") or sections.get("certifications", [])[:10],
        "skills": skills,
        "extraction_method": method,
    }

    if save:
        out_dir = Path("output")
        out_dir.mkdir(exist_ok=True)
        out_file = out_dir / (Path(path).stem + ".json")
        out_file.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    return result


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Parse a resume into RADIX skill JSON")
    ap.add_argument("file", help="path to a PDF / DOCX / TXT resume")
    ap.add_argument("--no-llm", action="store_true", help="use keyword mode only")
    args = ap.parse_args()

    data = parse_resume(args.file, use_llm=False if args.no_llm else None)
    print(json.dumps(data, indent=2, ensure_ascii=False))
    print(f"\nSaved to output/{Path(args.file).stem}.json  (method: {data['extraction_method']})")
    
